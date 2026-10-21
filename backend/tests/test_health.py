import asyncio
import io
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from openpyxl import load_workbook

from app.ai import service
from tests.conftest import RULE
from tests.test_ai import FakeProvider
from tests.test_api import client, server, wait_run  # noqa: F401

BROKEN = {**RULE, "list": {**RULE["list"], "item_selector": ".gone"}}


@pytest.fixture
def webhook():
    received: list[dict] = []

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            received.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b'{"errcode":0}')

    srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}/hook", received
    srv.shutdown()


async def crawl_once(client, source_id):
    run = (await client.post(f"/api/sources/{source_id}/crawl")).json()
    return await wait_run(client, run["id"])


async def test_health_alert_and_auto_repair(client, server, webhook, monkeypatch):
    url, received = webhook
    assert (await client.post("/api/settings/alert/test", json={"webhook_type": "wecom", "webhook_url": url})).json()["ok"]
    assert received.pop()["msgtype"] == "text"
    await client.put("/api/settings/alert", json={"failure_threshold": 2, "webhook_type": "wecom", "webhook_url": url})

    src = (await client.post("/api/sources", json={"name": "会改版的站", "start_url": f"{server}/list", "rate_limit": 20, "rule": BROKEN})).json()
    r1 = await crawl_once(client, src["id"])
    assert r1["status"] == "failed"
    assert any(e["msg"].startswith("页面片段") and "lot-item" in e["msg"] for e in r1["log"])
    s = (await client.get(f"/api/sources/{src['id']}")).json()
    assert s["health"] == "warning" and not received

    # AI 未开启：第二次失败判定为失效并告警，但不自动修复
    await crawl_once(client, src["id"])
    s = (await client.get(f"/api/sources/{src['id']}")).json()
    assert s["health"] == "broken" and "连续 2 次" in s["health_message"]
    assert "规则失效" in received[-1]["text"]["content"]

    # 开启 AI 自动修复：恢复 ok 后再次失效时，AI 重新生成规则并直接启用
    provider = FakeProvider([{"explanation": "修复", "rule": RULE}])
    monkeypatch.setattr(service, "provider_factory", lambda cfg: provider)
    await client.put("/api/ai/settings", json={"enabled": True, "daily_limit": 0, "auto_repair": True, "max_rounds": 2})
    await client.post("/api/ai/providers", json={"name": "c", "protocol": "anthropic", "api_key": "sk-1234567890ab", "model": "m"})
    await client.put(f"/api/sources/{src['id']}", json={"name": "会改版的站", "start_url": f"{server}/list", "rate_limit": 20, "rule": RULE})
    assert (await client.get(f"/api/sources/{src['id']}")).json()["health"] == "ok"
    await client.post(f"/api/sources/{src['id']}/rules", json={"config": BROKEN, "note": "站点改版"})
    await crawl_once(client, src["id"])
    await crawl_once(client, src["id"])

    # 修复在采集任务结束后异步进行
    for _ in range(50):
        rules = (await client.get(f"/api/sources/{src['id']}/rules")).json()
        if rules[0]["origin"] == "ai":
            break
        await asyncio.sleep(0.1)
    assert rules[0]["origin"] == "ai" and rules[0]["is_active"] and rules[0]["note"].startswith("规则失效")
    gens = (await client.get("/api/ai/generations")).json()["items"]
    assert gens[0]["trigger"] == "repair" and gens[0]["status"] == "success"
    assert ".gone" in provider.calls[0][0].content  # 提示中带上了失效的旧规则
    assert any("AI 已自动修复" in m["text"]["content"] for m in received)

    # 修复后的规则可以正常采集，健康状态恢复
    r = await crawl_once(client, src["id"])
    assert r["status"] == "success"
    assert (await client.get(f"/api/sources/{src['id']}")).json()["health"] == "ok"


async def test_export(client, server):
    src = (await client.post("/api/sources", json={"name": "导出站", "start_url": f"{server}/list", "rate_limit": 20, "rule": RULE})).json()
    await crawl_once(client, src["id"])

    r = await client.get("/api/lots/export", params={"format": "csv", "status": "live"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0].startswith("站点,标题") and "lot_no" in lines[0]
    assert len(lines) == 2 and "清乾隆" in lines[1] and "进行中" in lines[1]

    r = await client.get("/api/lots/export", params={"format": "xlsx", "sort": "-current_price"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 4 and ws["B2"].value == "民国 粉彩人物盘" and ws["D2"].value == 32000
