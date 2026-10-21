import csv
import io
from datetime import datetime
from typing import Any, Literal
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, HTTPException, Query, status
from fastapi.responses import Response
from sqlalchemy import Select, func, or_, select

from app.config import get_settings
from app.api.deps import SessionDep
from app.models import Lot, PriceSnapshot, Source
from app.schemas.api import LotDetailOut, LotOut, Page

router = APIRouter(prefix="/lots", tags=["lots"])

SortField = Literal[
    "last_seen_at", "-last_seen_at", "current_price", "-current_price",
    "end_time", "-end_time", "bid_count", "-bid_count", "first_seen_at", "-first_seen_at",
]


def _filtered(
    stmt: Select, source_id: int | None, status_: str | None, keyword: str | None,
    min_price: float | None, max_price: float | None,
) -> Select:
    if source_id:
        stmt = stmt.where(Lot.source_id == source_id)
    if status_:
        stmt = stmt.where(Lot.status == status_)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(Lot.title.ilike(like), Lot.unique_key.ilike(like)))
    price = func.coalesce(Lot.deal_price, Lot.current_price)
    if min_price is not None:
        stmt = stmt.where(price >= min_price)
    if max_price is not None:
        stmt = stmt.where(price <= max_price)
    return stmt


@router.get("", response_model=Page[LotOut])
async def list_lots(
    session: SessionDep,
    source_id: int | None = None,
    status_: str | None = Query(None, alias="status"),
    keyword: str | None = None,
    min_price: float | None = None,
    max_price: float | None = None,
    sort: SortField = "-last_seen_at",
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
) -> Page[LotOut]:
    args = (source_id, status_, keyword, min_price, max_price)
    total = await session.scalar(_filtered(select(func.count()).select_from(Lot), *args))
    column = getattr(Lot, sort.lstrip("-"))
    order = column.desc().nulls_last() if sort.startswith("-") else column.asc().nulls_last()
    stmt = (
        _filtered(select(Lot, Source.name).join(Source, Source.id == Lot.source_id), *args)
        .order_by(order, Lot.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    rows = (await session.execute(stmt)).all()
    items = [LotOut.model_validate(lot).model_copy(update={"source_name": name}) for lot, name in rows]
    return Page[LotOut](total=total or 0, items=items)


STATUS_LABEL = {"upcoming": "即将开始", "live": "进行中", "ended": "已结束", "unsold": "流拍"}
EXPORT_COLUMNS = [
    ("source_name", "站点"), ("title", "标题"), ("unique_key", "唯一标识"), ("current_price", "当前价"),
    ("start_price", "起拍价"), ("deal_price", "成交价"), ("currency", "币种"), ("bid_count", "出价次数"),
    ("status", "状态"), ("start_time", "开始时间"), ("end_time", "结束时间"), ("url", "链接"), ("image", "图片"),
    ("first_seen_at", "首次采集"), ("last_seen_at", "最近采集"),
]
EXPORT_LIMIT = 100_000


def _cell(value: Any, tz: ZoneInfo) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S")
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return str(value)
    return value


@router.get("/export")
async def export_lots(
    session: SessionDep,
    format: Literal["csv", "xlsx"] = "xlsx",
    source_id: int | None = None,
    status_: str | None = Query(None, alias="status"),
    keyword: str | None = None,
    min_price: float | None = None,
    max_price: float | None = None,
    sort: SortField = "-last_seen_at",
) -> Response:
    """按当前筛选条件导出拍品，扩展字段展开为独立列。"""
    column = getattr(Lot, sort.lstrip("-"))
    order = column.desc().nulls_last() if sort.startswith("-") else column.asc().nulls_last()
    stmt = (
        _filtered(select(Lot, Source.name).join(Source, Source.id == Lot.source_id), source_id, status_, keyword, min_price, max_price)
        .order_by(order, Lot.id.desc())
        .limit(EXPORT_LIMIT)
    )
    rows = (await session.execute(stmt)).all()
    tz = ZoneInfo(get_settings().timezone)
    extra_keys: list[str] = []
    for lot, _ in rows:
        for k in lot.extra or {}:
            if k not in extra_keys:
                extra_keys.append(k)
    header = [label for _, label in EXPORT_COLUMNS] + extra_keys
    table = []
    for lot, name in rows:
        values = {**{c: getattr(lot, c, None) for c, _ in EXPORT_COLUMNS}, "source_name": name}
        values["status"] = STATUS_LABEL.get(lot.status or "", lot.status)
        row = [_cell(values[c], tz) for c, _ in EXPORT_COLUMNS]
        row += [_cell((lot.extra or {}).get(k), tz) for k in extra_keys]
        table.append(row)

    filename = f"拍品数据_{datetime.now(tz).strftime('%Y%m%d_%H%M')}.{format}"
    disposition = f"attachment; filename*=UTF-8''{quote(filename)}"
    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(header)
        writer.writerows(table)
        # 带 BOM，Excel 直接打开不乱码
        return Response(("\ufeff" + buf.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": disposition})

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "拍品"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in table:
        ws.append([float(v) if hasattr(v, "is_finite") else v for v in row])
    ws.freeze_panes = "A2"
    for i, width in enumerate([14, 40, 24, 12, 12, 12, 8, 10, 10, 20, 20, 40, 30, 20, 20], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": disposition})


@router.get("/{lot_id}", response_model=LotDetailOut)
async def get_lot(lot_id: int, session: SessionDep) -> LotDetailOut:
    lot = await session.get(Lot, lot_id)
    if lot is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "拍品不存在")
    source = await session.get(Source, lot.source_id)
    snapshots = (
        await session.scalars(select(PriceSnapshot).where(PriceSnapshot.lot_id == lot_id).order_by(PriceSnapshot.captured_at))
    ).all()
    base = LotOut.model_validate(lot).model_dump()
    return LotDetailOut(**{**base, "source_name": source.name if source else None}, snapshots=snapshots)
